
# This piece of Python code was constructed to alter an excel spreadsheet. I created a new colomn with altered values as well as this, I made a barchat.


import openpyxl as xl
from openpyxl.chart import BarChart, Reference

wb = xl.load_workbook('transactions.xlsx')
sheet = wb['Sheet1']

for row in range(2, sheet.max_row + 1):
    cell = sheet.cell(row, 3)
    corrected_price = cell.value * 0.9

    corrected_price_cell = sheet.cell(row, 4)
    corrected_price_cell.value = corrected_price

sheet.cell(row=1, column=4).value = "Corrected Price"

values = Reference(sheet,
                   min_row=2,
                   max_row=sheet.max_row,
                   min_col=4,
                   max_col=4)

chart = BarChart(type='col')
chart.add_data(values)
chart.title = "Corrected Price Chart"
chart.y_axis.title = 'Corrected Price'

sheet.add_chart(chart, 'e2')

legend = chart.legend
legend.remove()

wb.save('transcations2.xlsx')
